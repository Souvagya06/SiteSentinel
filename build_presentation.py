"""Build the visual, project-focused SiteSentinel presentation."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SiteSentinel_Presentation.pptx"
prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
INK=RGBColor(14,25,43); NAVY=RGBColor(22,43,67); TEAL=RGBColor(0,151,157)
MINT=RGBColor(174,230,211); LIME=RGBColor(213,240,109); CORAL=RGBColor(244,119,82)
PAPER=RGBColor(246,246,239); WHITE=RGBColor(255,255,255); GREY=RGBColor(93,108,119)
PALE=RGBColor(226,235,230); LINE=RGBColor(200,213,208)

def shape(slide,x,y,w,h,fill,kind=MSO_SHAPE.ROUNDED_RECTANGLE,line=None):
    a=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=fill
    if line is None: a.line.fill.background()
    else: a.line.color.rgb=line; a.line.width=Pt(1.1)
    return a

def txt(slide,x,y,w,h,text,size=16,color=INK,bold=False,align=None,font="Aptos",margin=0):
    a=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=a.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin)
    for i,line in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
        p.space_after=Pt(4)
        if align is not None: p.alignment=align
    return a

def base(title,kicker,num,dark=False):
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=INK if dark else PAPER
    main=WHITE if dark else INK; muted=MINT if dark else GREY
    txt(s,.62,.25,10,.22,kicker.upper(),10,CORAL,True)
    txt(s,.62,.57,11.7,.58,title,26,main,True)
    shape(s,.62,1.25,.62,.055,LIME,MSO_SHAPE.RECTANGLE)
    txt(s,12.03,.36,.65,.34,f"{num:02d}",11,muted,True,PP_ALIGN.RIGHT)
    txt(s,.62,7.15,9,.17,"SITESENTINEL  /  STUDENT PROJECT",8,muted,True)
    return s

def img(slide,relative,x,y,w,h):
    path=ROOT/relative
    if path.exists():
        with PILImage.open(path) as image:
            aspect=image.width/image.height
        frame=w/h
        if aspect>frame:
            draw_w=w; draw_h=w/aspect
        else:
            draw_h=h; draw_w=h*aspect
        slide.shapes.add_picture(str(path),Inches(x+(w-draw_w)/2),Inches(y+(h-draw_h)/2),
                                 width=Inches(draw_w),height=Inches(draw_h))

def capsule(slide,x,y,w,label,fill=TEAL,color=WHITE):
    shape(slide,x,y,w,.34,fill)
    txt(slide,x+.08,y+.035,w-.16,.22,label,9,color,True,PP_ALIGN.CENTER)

def note(slide,x,y,w,text,dark=False):
    txt(slide,x,y,w,.38,text,10,MINT if dark else GREY)

# 1 — Photo-led title. The cover image is generated illustrative artwork, not a prototype photo.
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=INK
img(s,Path("patent")/"sitesentinel-cover-generated.png",0,0,13.333,7.5)
shape(s,0,0,6.55,7.5,INK,MSO_SHAPE.RECTANGLE)
txt(s,.85,.72,5,.32,"PROJECT PRESENTATION  /  2026",11,LIME,True)
txt(s,.85,1.72,5.6,1.1,"Site\nSentinel",39,WHITE,True)
shape(s,.88,4.17,.74,.07,CORAL,MSO_SHAPE.RECTANGLE)
txt(s,.85,4.5,5.35,1.15,"A camera-based aid for safer site entry",23,WHITE,True)
txt(s,.85,6.55,5.2,.35,"PPE checks  •  worker attendance  •  site dashboard",11,MINT,True)

# 2 — Problem, dramatized with concise issue → consequence → need flow.
s=base("The gap between site rules and site visibility","01  /  Problem statement",2)
blocks=[("PPE checks are intermittent","Supervisors cannot watch every worker and work area continuously."),
        ("Attendance is a separate task","Gate logs and safety observations may not share a worker-level record."),
        ("Small teams need practical tools","A dedicated safety team or enterprise platform may be beyond a small site’s reach.")]
for i,(h,b) in enumerate(blocks):
    x=.7+i*4.12
    shape(s,x,1.72,3.75,2.45,WHITE, line=LINE)
    shape(s,x,1.72,.12,2.45,[CORAL,TEAL,INK][i],MSO_SHAPE.RECTANGLE)
    txt(s,x+.3,2.05,3.12,.72,f"0{i+1}  {h}",18,INK,True)
    txt(s,x+.3,2.98,3.1,.88,b,14,GREY)
    if i<2: txt(s,x+3.82,2.72,.28,.5,"›",25,TEAL,True)
shape(s,.7,4.62,11.95,1.55,NAVY)
txt(s,1.02,4.9,2.0,.3,"THE NEED",11,LIME,True)
txt(s,1.02,5.3,10.95,.62,"Help a supervisor spot a PPE issue and connect a gate event to a worker record.",20,WHITE,True)
note(s,.72,6.4,11.3,"SiteSentinel supports supervision; it does not replace safety procedures or human judgement.")

# 3 — Proposed solution, one distinctive left-to-right process graphic.
s=base("One gate event, from camera to supervisor","02  /  Proposed solution",3,True)
steps=[("01","CAPTURE","Pi camera / stream"),("02","ANALYSE","YOLOv8 PPE + person detection"),
       ("03","IDENTIFY","Face match + attendance event"),("04","ACT","Dashboard + optional Pi alert")]
for i,(n,h,b) in enumerate(steps):
    x=.7+i*3.08
    shape(s,x,2.25,2.58,2.22,NAVY,line=TEAL)
    shape(s,x+.22,2.53,.48,.48,LIME,MSO_SHAPE.OVAL)
    txt(s,x+.22,2.63,.48,.21,n,10,INK,True,PP_ALIGN.CENTER)
    txt(s,x+.22,3.2,2.1,.3,h,13,LIME,True)
    txt(s,x+.22,3.63,2.12,.58,b,13,WHITE)
    if i<3: txt(s,x+2.68,3.08,.28,.48,"→",20,CORAL,True)
shape(s,.72,5.25,11.9,.83,RGBColor(33,57,80))
txt(s,1.0,5.49,11.25,.35,"Registered worker  →  PPE check  →  check-in / check-out log  →  supervisor view",15,WHITE,True,PP_ALIGN.CENTER)
note(s,.76,6.42,11.4,"PPE detection and worker association are prototype capabilities and still require validation across site conditions.",True)

# 4 — Innovation and competitors, with a compact comparison matrix.
s=base("A focused prototype in a broader safety market","03  /  Uniqueness and competitor analysis",4)
# table header and rows are native editable shapes
cols=[.72,4.0,7.02,9.93]; widths=[3.1,2.9,2.8,2.6]
shape(s,.72,1.65,11.9,.54,INK,MSO_SHAPE.RECTANGLE)
for x,w,t in zip(cols,widths,["APPROACH","WHAT IT OFFERS","RELATIVE SCOPE","POSITIONING"]): txt(s,x+.1,1.81,w-.2,.2,t,9,WHITE,True)
rows=[("Manual CCTV + rounds","Human-led observation","No automated detection","Baseline workflow"),
      ("Procore","Safety workflows and records","Construction management suite","Broader workflow platform"),
      ("viAct / Intenseye","AI video safety analytics","Commercial, broader analytics","Established AI platforms"),
      ("SiteSentinel","PPE + attendance + Pi feedback","Student prototype; gate-focused","Low-cost pilot hypothesis")]
for i,row in enumerate(rows):
    y=2.2+i*.77; fill=RGBColor(224,240,229) if i==3 else WHITE
    shape(s,.72,y,11.9,.66,fill,MSO_SHAPE.RECTANGLE,line=LINE)
    for x,w,t in zip(cols,widths,row): txt(s,x+.1,y+.14,w-.18,.42,t,12,TEAL if i==3 else INK,i==3)
shape(s,.72,5.63,11.9,.82,PALE)
txt(s,1.02,5.88,11.25,.31,"The distinction to test: one simple gate workflow joins safety checks with worker attendance and local alerts.",14,INK,True)
note(s,.75,6.58,11.5,"Scope references: Procore Quality & Safety; viAct PPE detection; Intenseye Safety. Product scope can vary by deployment.")

# 5 — Customers, roles and realistic initial market wedge.
s=base("Begin with one gate and one site team","04  /  Target users and market potential",5)
img(s,Path("patent")/"hardware setup.jpeg",7.55,1.65,5.0,2.55)
capsule(s,.78,1.7,1.4,"FIRST PILOT",TEAL)
txt(s,.78,2.28,6.15,.86,"Small and mid-sized contractors",24,INK,True)
txt(s,.78,3.2,6.05,.68,"A fixed entrance, an on-site supervisor and a manageable worker roster make the clearest starting point.",15,GREY)
roles=[("USER","Site safety supervisor"),("BUYER","Contractor / project manager"),("PARTNER","Site team + IT / camera support")]
for i,(tag,value) in enumerate(roles):
    y=4.58+i*.5; capsule(s,.8,y,1.12,tag,[CORAL,TEAL,INK][i]); txt(s,2.12,y+.03,4.75,.28,value,13,INK,True)
shape(s,7.55,4.55,5.0,1.42,NAVY)
txt(s,7.88,4.82,4.3,.25,"MARKET POTENTIAL",10,LIME,True)
txt(s,7.88,5.2,4.2,.55,"Validate demand in pilots before sizing the market.",17,WHITE,True)
note(s,.8,6.25,11.7,"Adoption depends on camera placement, privacy safeguards, worker consent, network access and site-specific PPE rules.")

# 6 — Budget displayed as two clear ranges rather than an unsupported cost pie chart.
s=base("Pilot investment and what it enables","05  /  Proposed budget and requirements",6,True)
txt(s,.75,1.62,5.6,.32,"INDICATIVE PLANNING RANGES  /  INR",10,LIME,True)
for x,label,value,detail,col in [
    (.78,"HARDWARE + SETUP","₹35k–₹60k","Pi, camera, enclosure, display, power + setup",TEAL),
    (6.78,"DEVELOPMENT EFFORT","₹1.5L–₹3L","Team effort for integration, calibration + pilot",CORAL),
]:
    shape(s,x,2.05,5.72,2.15,RGBColor(27,49,71),line=RGBColor(52,78,96))
    shape(s,x+.23,2.29,.12,1.63,col,MSO_SHAPE.RECTANGLE)
    txt(s,x+.55,2.34,4.82,.26,label,10,LIME,True)
    txt(s,x+.55,2.83,4.84,.62,value,30,WHITE,True)
    txt(s,x+.55,3.57,4.82,.42,detail,11,MINT)
txt(s,.8,4.98,11.7,.32,"RESOURCES FOR A PILOT",10,LIME,True)
items=["Pi + camera + enclosure","Existing GPU host","Site partner + consented test users","Mentor review + field calibration"]
for i,t in enumerate(items):
    x=.8+(i%2)*5.8; y=5.45+(i//2)*.48
    shape(s,x,y,.18,.18,CORAL,MSO_SHAPE.OVAL); txt(s,x+.3,y-.03,5.2,.28,t,12,WHITE)
note(s,.8,6.65,11.6,"Estimates for discussion, not supplier quotes. Host PC excluded from hardware; cloud and ongoing support costs are additional.",True)

# 7 — business model and roadmap as a path, not another list of cards.
s=base("A pilot-led route to a sustainable product","06  /  Business model and future scope",7)
path=[("PILOT","One entrance\nMeasure fit"),("REFINE","Improve matching\nReview false alerts"),("OFFER","Setup + monthly\nsite subscription"),("EXTEND","More cameras\nRules + reporting")]
for i,(h,b) in enumerate(path):
    x=.8+i*3.05
    shape(s,x,2.15,2.48,1.82,[TEAL,NAVY,CORAL,INK][i])
    txt(s,x+.2,2.48,2.05,.3,f"0{i+1}  {h}",13,LIME if i!=2 else WHITE,True)
    txt(s,x+.2,3.02,2.05,.62,b,14,WHITE,True)
    if i<3: txt(s,x+2.57,2.85,.4,.38,"→",20,TEAL,True)
shape(s,.8,4.52,7.78,1.33,PALE)
txt(s,1.08,4.78,3.0,.25,"PRICE HYPOTHESIS",10,TEAL,True)
txt(s,1.08,5.15,6.9,.44,"₹2,000–₹5,000 / site / month  +  setup",17,INK,True)
shape(s,8.92,4.52,3.75,1.33,NAVY)
txt(s,9.2,4.78,3.15,.25,"TO VALIDATE",10,LIME,True)
txt(s,9.2,5.15,3.05,.44,"Support cost + willingness to pay",13,WHITE,True)
note(s,.82,6.35,11.7,"Price is a customer-discovery hypothesis, not an approved commercial offer. Expand scope only after pilot evidence.")

# 8 — evidence of work and an engaging, presenter-ready live demo flow.
s=base("What is built — and what to show live","07  /  Demonstration and progress",8)
img(s,Path("patent")/"prototype enclosure.jpeg",.72,1.58,4.15,3.14)
img(s,Path("patent")/"hardware setup.jpeg",5.05,1.58,4.15,3.14)
shape(s,9.38,1.58,3.2,3.14,NAVY)
txt(s,9.72,1.9,2.5,.27,"SOFTWARE",10,LIME,True)
txt(s,9.72,2.42,2.45,1.8,"Worker setup\nPPE detection\nAttendance log\nWeb dashboard\nPi alert control",14,WHITE,True)
txt(s,.8,4.85,3.95,.25,"PiBox enclosure prototype",11,GREY,True,PP_ALIGN.CENTER)
txt(s,5.12,4.85,4.0,.25,"Camera, Pi and alert circuit",11,GREY,True,PP_ALIGN.CENTER)
steps=["Register a test worker","Run a camera check","Show the log + dashboard","Trigger Pi feedback"]
for i,t in enumerate(steps):
    x=.8+i*3.05; shape(s,x,5.45,.4,.4,TEAL,MSO_SHAPE.OVAL); txt(s,x,5.54,.4,.18,str(i+1),10,WHITE,True,PP_ALIGN.CENTER)
    txt(s,x+.52,5.47,2.32,.43,t,12,INK,True)
note(s,.82,6.35,11.6,"Working prototype components are present; site-level accuracy, reliability and cost remain to be measured.")

prs.save(OUT)
print(f"Saved {OUT} ({len(prs.slides)} slides)")
